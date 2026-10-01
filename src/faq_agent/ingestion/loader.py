import hashlib
from dataclasses import replace
from pathlib import Path

from faq_agent.schemas import SourceDocument

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


def load_document(path: Path) -> SourceDocument:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {path.suffix}")

    if suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8")
    elif suffix == ".pdf":
        text = _read_pdf(path)
    elif suffix == ".docx":
        text = _read_docx(path)
    else:
        raise ValueError(f"Unsupported file type: {path.suffix}")

    clean_text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not clean_text:
        raise ValueError(f"No readable text found in {path}")

    return SourceDocument(
        source_id=path.stem.lower().replace(" ", "-"),
        path=str(path),
        title=path.stem,
        text=clean_text,
        metadata={"extension": suffix, "filename": path.name},
    )


def load_documents(folder: Path) -> list[SourceDocument]:
    if not folder.is_dir():
        raise ValueError("Source directory does not exist")
    docs: list[SourceDocument] = []
    for path in sorted(folder.glob("**/*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            if path.is_symlink() or not path.resolve().is_relative_to(folder.resolve()):
                raise ValueError("Source links outside the corpus are not allowed")
            relative = path.relative_to(folder).as_posix()
            source_id = hashlib.sha256(relative.encode()).hexdigest()[:24]
            docs.append(replace(load_document(path), source_id=source_id))
    return docs


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)
